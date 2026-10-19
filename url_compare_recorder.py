#!/usr/bin/env python3
"""Two-application URL performance comparison recorder for Microsoft Edge (Python 3.9).

Requires: Windows, Microsoft Edge, openpyxl (pip install openpyxl).
No Selenium, Playwright, WebDriver, or application-code changes.

Use:
    py -3.9 url_compare_recorder.py --app1 "https://old.example.local/" --name1 "CRM Old" ^
                                    --app2 "https://new.example.local/" --name2 "CRM New"

What it does:
  * Opens the two applications in two separate tabs of a private Edge window.
    Work in whichever tab you like - each tab is recorded on its own.
    Tabs you open yourself (Ctrl+T, popups) are picked up automatically.
  * Every click / page navigation is timed per tab.
  * Actions done in one application but not in the other are flagged in the
    report ("Not performed in CRM New").
  * Stop with Ctrl+C in this console (or close the Edge window). A formatted
    Excel report url_compare_report_<timestamp>.xlsx is written.
"""

import argparse
import base64
import datetime as dt
import hashlib
import json
import os
import select
import shutil
import statistics
import socket
import struct
import subprocess
import sys
import tempfile
import time
import urllib.request
import urllib.error
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None


# A click is complete when its tracked network requests finish and there have
# been no new tracked requests for QUIET_MS. This is NOT exact visual readiness.
QUIET_MS = 1200
NO_REQUEST_WAIT_MS = 2000
MAX_ACTION_MS = 120000
MAX_WS_FRAME_BYTES = 16 * 1024 * 1024
TAB_POLL_SECONDS = 2.0

DEFAULT_SLOW_MS = 3000       # actions at/above this are highlighted as slow
SIMILAR_MS = 200             # differences under max(SIMILAR_MS, 10%) are "Similar"

# Change to True if your application uses custom clickable divs rather than
# normal HTML buttons, links, or accessible role=button controls.
RECORD_ALL_CLICKS = False

INJECT_JS = r"""
(function () {
  if (window.__edgePerfRecorderInstalled) return;
  window.__edgePerfRecorderInstalled = true;
  var allClicks = __ALL_CLICKS__;
  document.addEventListener('click', function(ev) {
    try {
      var target = ev.target;
      if (!target || !target.closest) return;
      var hit = target.closest('button, a, input[type="submit"], input[type="button"], input[type="image"], [role="button"], [role="tab"], [onclick], [tabindex]');
      if (!hit && !allClicks) return;
      var el = hit || target;
      var raw = el.getAttribute('aria-label') || el.innerText || el.value ||
                el.title || el.id || el.getAttribute('name') || el.tagName;
      var label = String(raw || 'Unnamed control').replace(/\s+/g, ' ').trim().slice(0, 100);
      var url = location.origin + location.pathname;
      window.__edgePerfNotify(JSON.stringify({ label: label, url: url, timeMs: Date.now() }));
    } catch (ex) { /* Don't interfere with application behavior. */ }
  }, true);
})();
""".replace('__ALL_CLICKS__', 'true' if RECORD_ALL_CLICKS else 'false')


def clean_url(raw):
    """Never record query strings or URL fragments (potentially sensitive)."""
    try:
        parts = urlsplit(raw or '')
        return urlunsplit((parts.scheme, parts.netloc, parts.path, '', ''))
    except ValueError:
        return '[invalid URL]'


def short_url(url):
    path = urlsplit(url).path.rstrip('/')
    return path.rsplit('/', 1)[-1] or url


def median_or_none(values):
    return round(statistics.median(values)) if values else None


def find_edge():
    paths = [
        shutil.which('msedge'),
        os.path.join(os.environ.get('PROGRAMFILES(X86)', ''), 'Microsoft', 'Edge', 'Application', 'msedge.exe'),
        os.path.join(os.environ.get('PROGRAMFILES', ''), 'Microsoft', 'Edge', 'Application', 'msedge.exe'),
        os.path.join(os.environ.get('LOCALAPPDATA', ''), 'Microsoft', 'Edge', 'Application', 'msedge.exe'),
    ]
    for path in paths:
        if path and os.path.isfile(path):
            return path
    raise RuntimeError('Microsoft Edge (msedge.exe) was not found on this computer.')


def wait_for_debugging_port(profile, process, seconds=25):
    port_file = Path(profile) / 'DevToolsActivePort'
    deadline = time.monotonic() + seconds
    while time.monotonic() < deadline:
        if port_file.exists():
            lines = port_file.read_text(encoding='utf-8', errors='replace').splitlines()
            if lines and lines[0].isdigit():
                return int(lines[0])
        if process.poll() is not None:
            raise RuntimeError('Edge exited before it opened the debugging endpoint.')
        time.sleep(0.2)
    raise RuntimeError('Edge did not enable local debugging. IT browser policy may block it.')


class WSClient:
    """Minimal RFC 6455 WebSocket client for localhost CDP, stdlib only."""

    def __init__(self, ws_url):
        parsed = urlsplit(ws_url)
        if parsed.scheme != 'ws' or parsed.hostname not in ('localhost', '127.0.0.1', '::1'):
            raise RuntimeError('Refusing to connect to a non-local debugging WebSocket.')
        host = parsed.hostname
        port = parsed.port or 80
        path = parsed.path + ('?' + parsed.query if parsed.query else '')
        self.sock = socket.create_connection((host, port), timeout=10)
        self.sock.settimeout(10)
        key = base64.b64encode(os.urandom(16)).decode('ascii')
        request = (
            'GET %s HTTP/1.1\r\nHost: %s:%d\r\nUpgrade: websocket\r\n'
            'Connection: Upgrade\r\nSec-WebSocket-Version: 13\r\n'
            'Sec-WebSocket-Key: %s\r\n\r\n'
        ) % (path, host, port, key)
        self.sock.sendall(request.encode('ascii'))
        response = bytearray()
        while b'\r\n\r\n' not in response:
            block = self.sock.recv(4096)
            if not block:
                raise ConnectionError('Browser closed during WebSocket handshake.')
            response.extend(block)
            if len(response) > 65536:
                raise RuntimeError('Oversized WebSocket handshake.')
        hdrs, leftovers = bytes(response).split(b'\r\n\r\n', 1)
        expected = base64.b64encode(hashlib.sha1((key + '258EAFA5-E914-47DA-95CA-C5AB0DC85B11').encode('ascii')).digest()).decode('ascii')
        if not hdrs.startswith(b'HTTP/1.1 101') or ('sec-websocket-accept: ' + expected).lower().encode() not in hdrs.lower():
            raise RuntimeError('CDP WebSocket handshake was rejected: %r' % hdrs[:250])
        self.buffer = bytearray(leftovers)
        self.sock.settimeout(15)

    def _exact(self, n):
        while len(self.buffer) < n:
            chunk = self.sock.recv(max(4096, n - len(self.buffer)))
            if not chunk:
                raise ConnectionError('Edge browser connection closed.')
            self.buffer.extend(chunk)
        out = bytes(self.buffer[:n])
        del self.buffer[:n]
        return out

    def send_frame(self, opcode, payload):
        # Browser clients MUST mask WebSocket frames.
        mask = os.urandom(4)
        n = len(payload)
        if n < 126:
            length_part = bytes((0x80 | n,))
        elif n < 65536:
            length_part = bytes((0x80 | 126,)) + struct.pack('!H', n)
        else:
            length_part = bytes((0x80 | 127,)) + struct.pack('!Q', n)
        payload = bytes(byte ^ mask[i % 4] for i, byte in enumerate(payload))
        self.sock.sendall(bytes((0x80 | opcode,)) + length_part + mask + payload)

    def send_json(self, obj):
        self.send_frame(1, json.dumps(obj, separators=(',', ':')).encode('utf-8'))

    def recv_json(self, wait_seconds=0.2):
        if not self.buffer:
            readable, _, _ = select.select([self.sock], [], [], wait_seconds)
            if not readable:
                return None
        message = bytearray()
        started = False
        while True:
            b1, b2 = self._exact(2)
            opcode, fin = b1 & 0x0f, bool(b1 & 0x80)
            size = b2 & 0x7f
            if size == 126:
                size = struct.unpack('!H', self._exact(2))[0]
            elif size == 127:
                size = struct.unpack('!Q', self._exact(8))[0]
            if size > MAX_WS_FRAME_BYTES:
                raise RuntimeError('WebSocket message exceeds safety limit.')
            mask = self._exact(4) if b2 & 0x80 else None
            payload = self._exact(size)
            if mask:
                payload = bytes(byte ^ mask[i % 4] for i, byte in enumerate(payload))
            if opcode == 8:
                raise ConnectionError('Edge closed the debugging WebSocket.')
            if opcode == 9:
                self.send_frame(10, payload)
                continue
            if opcode == 10:
                continue
            if opcode in (1, 2):
                started = True
            elif opcode != 0 or not started:
                continue
            message.extend(payload)
            if len(message) > MAX_WS_FRAME_BYTES:
                raise RuntimeError('WebSocket message exceeds safety limit.')
            if fin:
                return json.loads(message.decode('utf-8'))

    def close(self):
        try:
            self.send_frame(8, b'')
        except (OSError, ConnectionError):
            pass
        self.sock.close()


class Report:
    """Shared state for all tabs: application names/URLs, action numbering and results."""

    def __init__(self, app1, app2, name1, name2, slow_ms):
        self.name1, self.name2 = name1, name2
        self.app_urls = {name1: clean_url(app1), name2: clean_url(app2)}
        self.slow_ms = slow_ms
        self.seq = 0
        self.completed = []          # list of (result dict, request list)
        self.started = dt.datetime.now()

    def next_id(self):
        self.seq += 1
        return self.seq

    def app_name(self, url, default=None):
        # Match complete URL first, then the application origin (host + port),
        # then the application the tab was opened for (covers SSO redirects).
        cleaned = clean_url(url)
        exact = [(len(base), name) for name, base in self.app_urls.items()
                 if base and cleaned.startswith(base)]
        if exact:
            return max(exact)[1]
        parsed = urlsplit(cleaned)
        origins = []
        for name, base in self.app_urls.items():
            if not base:
                continue
            ref = urlsplit(base)
            if ref.scheme == parsed.scheme and ref.netloc.lower() == parsed.netloc.lower():
                origins.append(name)
        if len(origins) == 1:
            return origins[0]
        return default or parsed.hostname or 'Unknown'

    def rel_path(self, url, app):
        """Page path relative to the application's base folder, so the same page
        in both applications gets the same key even if virtual directories differ."""
        path = urlsplit(url).path or '/'
        base = self.app_urls.get(app) or ''
        base_dir = urlsplit(base).path.rsplit('/', 1)[0] + '/' if base else '/'
        if base_dir != '/' and path.lower().startswith(base_dir.lower()):
            path = '/' + path[len(base_dir):]
        return path.lower()


class Tab:
    """One Edge tab: its own DevTools connection and its own in-flight action."""

    def __init__(self, ws, report, name, default_app=None):
        self.ws = ws
        self.report = report
        self.name = name
        self.default_app = default_app
        self.cmd_seq = 0
        self.commands = {}
        self.active = None
        self.reqs = {}
        self.main_frame = None
        self.cdp_offset_ms = None

    def cmd(self, name, params=None):
        self.cmd_seq += 1
        self.commands[self.cmd_seq] = name
        self.ws.send_json({'id': self.cmd_seq, 'method': name, 'params': params or {}})

    def start(self):
        self.cmd('Runtime.enable')
        self.cmd('Page.enable')
        self.cmd('Network.enable', {'maxTotalBufferSize': 10000000})
        self.cmd('Runtime.addBinding', {'name': '__edgePerfNotify'})
        self.cmd('Page.addScriptToEvaluateOnNewDocument', {'source': INJECT_JS})
        self.cmd('Runtime.evaluate', {'expression': INJECT_JS})
        self.cmd('Page.getFrameTree')

    def new_action(self, label, page, start_ms, kind='click'):
        if self.active:
            self.finish_action('INTERRUPTED_BY_NEXT_ACTION')
        app = self.report.app_name(page, self.default_app)
        rel = self.report.rel_path(page, app)
        if kind == 'navigation':
            label = 'Open page ' + rel
        action = dict(id=self.report.next_id(), kind=kind, app=app, page=clean_url(page),
                      rel=rel, label=label[:110], started_ms=start_ms, requests=[], pending=set(),
                      last_activity_ms=start_ms, last_end_ms=None, dom_ms=None, load_ms=None)
        self.active = action
        print('[%03d] %-7s %-16s %-35s %s' %
              (action['id'], self.name[:7], app[:16], action['label'][:35], rel[:70]), flush=True)
        return action

    def on_event(self, message):
        if 'id' in message:
            method = self.commands.pop(message['id'], 'unknown')
            if 'error' in message:
                print('DevTools error for %s: %s' % (method, message['error']), file=sys.stderr)
                if method in ('Runtime.addBinding', 'Page.addScriptToEvaluateOnNewDocument', 'Network.enable'):
                    raise RuntimeError('Required Edge DevTools feature not available: ' + method)
            if method == 'Page.getFrameTree':
                self.main_frame = message.get('result', {}).get('frameTree', {}).get('frame', {}).get('id')
            return
        method = message.get('method')
        p = message.get('params', {})
        if method == 'Page.frameNavigated':
            frame = p.get('frame', {})
            if not frame.get('parentId'):
                self.main_frame = frame.get('id', self.main_frame)
        elif method == 'Runtime.bindingCalled':
            if p.get('name') == '__edgePerfNotify':
                try:
                    body = json.loads(p['payload'])
                    self.new_action(body.get('label') or 'Unnamed click',
                                    body.get('url') or '', float(body.get('timeMs', time.time() * 1000)))
                except (ValueError, KeyError, TypeError) as ex:
                    print('Could not parse a click event: %s' % ex, file=sys.stderr)
        elif method == 'Network.requestWillBeSent':
            self.request_start(p)
        elif method == 'Network.responseReceived':
            req = self.reqs.get(p.get('requestId'))
            if req:
                self.apply_response(req, p.get('response') or {})
        elif method in ('Network.loadingFinished', 'Network.loadingFailed'):
            self.request_done(p, failed=(method == 'Network.loadingFailed'))
        elif method in ('Page.domContentEventFired', 'Page.loadEventFired'):
            a = self.active
            if a and self.cdp_offset_ms is not None and 'timestamp' in p:
                event_ms = float(p['timestamp']) * 1000 + self.cdp_offset_ms
                delta = round(event_ms - a['started_ms'])
                if delta >= 0:
                    a['dom_ms' if method == 'Page.domContentEventFired' else 'load_ms'] = delta

    @staticmethod
    def apply_response(req, response):
        status = response.get('status')
        req['status'] = int(status) if isinstance(status, (int, float)) else None
        timings = response.get('timing') or {}
        if 'receiveHeadersEnd' in timings and 'sendEnd' in timings:
            req['ttfb_ms'] = max(0, round(timings['receiveHeadersEnd'] - timings['sendEnd'], 1))

    def request_start(self, p):
        rid = p.get('requestId')
        if not rid:
            return
        ts = float(p.get('timestamp', 0))
        wall = float(p.get('wallTime', time.time())) * 1000
        if ts:
            self.cdp_offset_ms = wall - ts * 1000
        # A redirect re-uses the requestId: close the previous hop first.
        if p.get('redirectResponse') and rid in self.reqs:
            self.apply_response(self.reqs[rid], p['redirectResponse'])
            self.request_done({'requestId': rid, 'timestamp': ts})
        request = p.get('request') or {}
        url = clean_url(request.get('url', ''))
        if not url.startswith(('http://', 'https://')):
            return
        kind = p.get('type', '')
        if (kind == 'Document' and not self.active and
                (not self.main_frame or p.get('frameId') == self.main_frame)):
            self.new_action('', url, wall, 'navigation')
        action = self.active
        if not action or wall < action['started_ms'] - 150:
            return
        # The new request belongs to the active click/navigation.
        req = dict(id=rid, action=action, method=request.get('method', ''), url=url,
                   type=kind, started_ts=ts, started_ms=wall, status=None,
                   duration_ms=None, ttfb_ms=None, bytes_received=0, error='')
        self.reqs[rid] = req
        action['requests'].append(req)
        action['pending'].add(rid)
        action['last_activity_ms'] = max(action['last_activity_ms'], wall)

    def request_done(self, p, failed=False):
        rid = p.get('requestId')
        req = self.reqs.pop(rid, None)
        if not req:
            return
        end_ts = float(p.get('timestamp', req['started_ts']))
        elapsed = max(0, (end_ts - req['started_ts']) * 1000)
        req['duration_ms'] = round(elapsed, 1)
        req['bytes_received'] = p.get('encodedDataLength', 0) if not failed else 0
        if failed:
            req['error'] = p.get('errorText') or 'Network failed'
        action = req['action']
        action['pending'].discard(rid)
        end_wall = req['started_ms'] + elapsed
        action['last_end_ms'] = max(action['last_end_ms'] or 0, end_wall)
        action['last_activity_ms'] = max(action['last_activity_ms'], end_wall)

    def tick(self):
        action = self.active
        if not action:
            return
        now = time.time() * 1000
        if now - action['started_ms'] > MAX_ACTION_MS:
            self.finish_action('TIMEOUT')
        elif action['pending']:
            return
        elif action['requests'] and now - action['last_activity_ms'] >= QUIET_MS:
            self.finish_action('NETWORK_QUIET')
        elif not action['requests'] and now - action['started_ms'] >= NO_REQUEST_WAIT_MS:
            self.finish_action('NO_NETWORK_OBSERVED')

    def finish_action(self, reason):
        action = self.active
        if not action:
            return
        self.active = None
        # The duration here is from click to LAST network completion, not the
        # 1.2-second confirmation window used before declaring network quiet.
        total_ms = None
        if reason == 'NETWORK_QUIET' and action['last_end_ms'] is not None:
            total_ms = round(max(0, action['last_end_ms'] - action['started_ms']))
        ended = [r for r in action['requests'] if r['duration_ms'] is not None]
        slow = max(ended, key=lambda r: r['duration_ms']) if ended else None
        result = {
            'action_id': action['id'],
            'started': dt.datetime.fromtimestamp(action['started_ms'] / 1000),
            'tab': self.name, 'application': action['app'], 'page': action['page'],
            'page_rel': action['rel'], 'action': action['label'], 'action_type': action['kind'],
            'outcome': reason, 'total_ms': total_ms,
            'dom_ms': action['dom_ms'], 'load_ms': action['load_ms'],
            'request_count': len(action['requests']),
            'failed_requests': sum(1 for r in action['requests'] if r['error'] or
                                   (isinstance(r['status'], int) and r['status'] >= 400)),
            'slowest_ms': round(slow['duration_ms']) if slow else None,
            'slowest_url': slow['url'] if slow else '',
            'server_wait_ms': round(slow['ttfb_ms']) if slow and slow['ttfb_ms'] is not None else None,
            'kb': round(sum(float(r['bytes_received'] or 0) for r in ended) / 1024, 1),
            'match_key': action['rel'] + '|' + action['label'].strip().lower(),
        }
        self.report.completed.append((result, action['requests']))
        shown = '%d ms' % total_ms if total_ms is not None else 'not measured'
        print('      -> %-20s %-13s requests: %d' % (reason, shown, len(action['requests'])), flush=True)


class Browser:
    """Edge process + one Tab per open page. New tabs are attached automatically."""

    def __init__(self, port, report):
        self.port = port
        self.report = report
        self.opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
        self.tabs = []
        self.known = set()
        self.extra = 0

    def _get(self, path, method='GET'):
        # The endpoint is LOCAL to the user's machine. Some enterprise/Anaconda
        # environments configure HTTP_PROXY/HTTPS_PROXY or a Windows proxy which
        # routes ordinary urllib calls through a corporate gateway. That gateway
        # can return "HTTP Error 403: Forbidden" for 127.0.0.1 even though
        # Edge launched correctly. Bypass proxies ONLY for this loopback call.
        req = urllib.request.Request('http://127.0.0.1:%d%s' % (self.port, path), method=method)
        try:
            with self.opener.open(req, timeout=5) as res:
                return json.load(res)
        except urllib.error.HTTPError as ex:
            raise RuntimeError(
                'Edge LOCAL debugging endpoint returned HTTP %s. '
                'Check whether corporate Edge policy blocks remote debugging. '
                'This error is not from either application URL.' % ex.code
            ) from ex

    def pages(self):
        return [t for t in self._get('/json/list')
                if t.get('type') == 'page' and t.get('webSocketDebuggerUrl')
                and not str(t.get('url', '')).startswith('devtools://')]

    def new_tab(self):
        # Newer Edge requires PUT for /json/new; older versions accept GET.
        try:
            return self._get('/json/new?about:blank', method='PUT')
        except RuntimeError:
            return self._get('/json/new?about:blank')

    def attach(self, target, name, default_app=None, start_url=None):
        self.known.add(target['id'])
        tab = Tab(WSClient(target['webSocketDebuggerUrl']), self.report, name, default_app)
        tab.start()
        if start_url:
            # Navigate only after instrumentation is attached.
            tab.cmd('Page.navigate', {'url': start_url})
        self.tabs.append(tab)
        return tab

    def sync_tabs(self):
        """Attach to tabs the user opened (Ctrl+T, links opening a new window)."""
        try:
            targets = self.pages()
        except (RuntimeError, OSError, ValueError):
            return
        for target in targets:
            if target['id'] in self.known:
                continue
            self.extra += 1
            name = 'Tab %d' % (2 + self.extra)
            url = target.get('url', '')
            try:
                self.attach(target, name, self.report.app_name(url) if url.startswith('http') else None)
                print('Now also recording %s (%s)' % (name, clean_url(url)[:80]))
            except (RuntimeError, OSError, ConnectionError) as ex:
                print('Could not attach to a new tab: %s' % ex, file=sys.stderr)

    def drop(self, tab, why):
        if tab.active:
            tab.finish_action('TAB_CLOSED')
        try:
            tab.ws.close()
        except OSError:
            pass
        self.tabs.remove(tab)
        print('%s is no longer recorded (%s).' % (tab.name, why))

    def pump(self):
        """Read every waiting DevTools message from every tab, then tick timers."""
        ready = [t for t in self.tabs if t.ws.buffer]
        if not ready and self.tabs:
            readable, _, _ = select.select([t.ws.sock for t in self.tabs], [], [], 0.2)
            ready = [t for t in self.tabs if t.ws.sock in readable]
        elif not self.tabs:
            time.sleep(0.2)
        for tab in ready:
            try:
                for _ in range(500):
                    event = tab.ws.recv_json(0)
                    if event is None:
                        break
                    tab.on_event(event)
            except (ConnectionError, OSError, RuntimeError, ValueError) as ex:
                self.drop(tab, ex)
        for tab in self.tabs:
            tab.tick()


# ----------------------------------------------------------------------------
# Excel report
# ----------------------------------------------------------------------------

NAVY = '1F3864'
BLUE = '2F5496'
ZEBRA = 'F2F5FA'
GREY_TEXT = '595959'
THIN = Side(style='thin', color='BFBFBF') if Workbook else None


def _fill(color):
    return PatternFill('solid', start_color=color, end_color=color)


def _title(ws, text, subtitle, width_cols):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width_cols)
    cell = ws.cell(row=1, column=1, value=text)
    cell.font = Font(size=16, bold=True, color='FFFFFF')
    cell.fill = _fill(NAVY)
    cell.alignment = Alignment(vertical='center', indent=1)
    for c in range(2, width_cols + 1):
        ws.cell(row=1, column=c).fill = _fill(NAVY)
    ws.row_dimensions[1].height = 32
    sub = ws.cell(row=2, column=1, value=subtitle)
    sub.font = Font(italic=True, color=GREY_TEXT)


def _table(ws, row, headers, rows, number_cols=()):
    """Write a styled table starting at `row`. Returns the last row written."""
    for c, head in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=head)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = _fill(BLUE)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    ws.row_dimensions[row].height = 30
    for i, values in enumerate(rows, 1):
        for c, value in enumerate(values, 1):
            cell = ws.cell(row=row + i, column=c, value=value)
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            cell.alignment = Alignment(vertical='top', wrap_text=isinstance(value, str) and len(value) > 40)
            if i % 2 == 0:
                cell.fill = _fill(ZEBRA)
            if c in number_cols and isinstance(value, (int, float)):
                cell.number_format = '#,##0'
            if isinstance(value, dt.datetime):
                cell.number_format = 'yyyy-mm-dd hh:mm:ss'
    return row + len(rows)


def _widths(ws, widths):
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w


def _section(ws, row, text, cols=5):
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(size=12, bold=True, color=NAVY)
    for c in range(1, cols + 1):
        ws.cell(row=row, column=c).border = Border(bottom=Side(style='medium', color=NAVY))
    return row + 1


def _data_sheet(wb, name, title, subtitle, headers, rows, widths, number_cols=()):
    ws = wb.create_sheet(name)
    _title(ws, title, subtitle, len(headers))
    last = _table(ws, 3, headers, rows, number_cols)
    ws.freeze_panes = 'A4'
    ws.auto_filter.ref = 'A3:%s%d' % (get_column_letter(len(headers)), max(last, 3))
    _widths(ws, widths)
    return ws


def build_comparison(report):
    """Group actions by page + label and compare application 1 with application 2."""
    n1, n2 = report.name1, report.name2
    groups = {}
    for result, _ in report.completed:
        if result['application'] not in (n1, n2):
            continue
        g = groups.setdefault(result['match_key'], {
            'action': result['action'], 'page': result['page_rel'], n1: [], n2: []})
        g[result['application']].append(result)
    rows = []
    for g in groups.values():
        r1, r2 = g[n1], g[n2]
        med1 = median_or_none([r['total_ms'] for r in r1 if r['total_ms'] is not None])
        med2 = median_or_none([r['total_ms'] for r in r2 if r['total_ms'] is not None])
        diff = pct = None
        verdict = ''
        if med1 is not None and med2 is not None:
            diff = med2 - med1
            pct = (diff / med1) if med1 else None
            if abs(diff) < max(SIMILAR_MS, 0.1 * min(med1, med2)):
                verdict = 'Similar'
            else:
                verdict = '%s slower' % (n2 if diff > 0 else n1)
        if r1 and r2:
            status = 'Performed in both'
            note = '' if diff is not None else 'Timing not captured cleanly in one application'
        elif r1:
            status, note = 'Only in %s' % n1, 'Not performed in %s - no comparison possible' % n2
        else:
            status, note = 'Only in %s' % n2, 'Not performed in %s - no comparison possible' % n1
        rows.append(dict(action=g['action'], page=g['page'], runs1=len(r1), med1=med1, runs2=len(r2),
                         med2=med2, diff=diff, pct=pct, verdict=verdict, status=status, note=note))
    return rows


def write_excel(report, path):
    wb = Workbook()
    n1, n2 = report.name1, report.name2
    slow_ms = report.slow_ms
    results = [r for r, _ in report.completed]
    comparison = build_comparison(report)
    in_both = {(c['page'], c['action'].strip().lower()) for c in comparison if c['runs1'] and c['runs2']}
    finished = dt.datetime.now()
    subtitle = 'Generated %s   |   Slow threshold: %s ms   |   Times are click to last network response' % (
        finished.strftime('%d %b %Y %H:%M'), '{:,}'.format(slow_ms))

    # ---- Summary -----------------------------------------------------------
    ws = wb.active
    ws.title = 'Summary'
    _title(ws, '%s vs %s - Performance Comparison' % (n1, n2), subtitle, 5)
    _widths(ws, [44, 24, 24, 4, 60])
    row = _section(ws, 4, 'Test details')
    details = [
        (n1 + ' URL', report.app_urls[n1] or '(not given)'),
        (n2 + ' URL', report.app_urls[n2] or '(not given)'),
        ('Recording started', report.started),
        ('Recording ended', finished),
        ('Duration (minutes)', round((finished - report.started).total_seconds() / 60, 1)),
    ]
    for label, value in details:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        cell = ws.cell(row=row, column=2, value=value)
        if isinstance(value, dt.datetime):
            cell.number_format = 'yyyy-mm-dd hh:mm:ss'
            cell.alignment = Alignment(horizontal='left')
        row += 1

    row = _section(ws, row + 1, 'Results at a glance')

    def per_app(app):
        rs = [r for r in results if r['application'] == app]
        timed = [r['total_ms'] for r in rs if r['total_ms'] is not None]
        only = sum(1 for c in comparison if c['status'] == 'Only in %s' % app)
        return [len(rs), len(timed), median_or_none(timed), max(timed) if timed else None,
                sum(1 for t in timed if t >= slow_ms), sum(r['failed_requests'] for r in rs), only]

    a1, a2 = per_app(n1), per_app(n2)
    labels = ['Actions recorded', 'Actions timed', 'Median action time (ms)', 'Slowest action (ms)',
              'Slow actions (>= %s ms)' % '{:,}'.format(slow_ms), 'Failed requests (errors / HTTP 4xx-5xx)',
              'Actions not performed in the other application']
    row = _table(ws, row, ['Metric', n1, n2], [[l, x, y] for l, x, y in zip(labels, a1, a2)], (2, 3))
    for r in range(row - len(labels) + 1, row + 1):
        x, y = ws.cell(row=r, column=2).value, ws.cell(row=r, column=3).value
        if r - (row - len(labels)) in (3, 4) and isinstance(x, int) and isinstance(y, int) and x != y:
            worse = 2 if x > y else 3
            ws.cell(row=r, column=worse).font = Font(bold=True, color='C00000')

    compared = [c for c in comparison if c['verdict']]
    row += 2
    if compared:
        s2 = sum(1 for c in compared if c['verdict'] == n2 + ' slower')
        s1 = sum(1 for c in compared if c['verdict'] == n1 + ' slower')
        verdict = ('Verdict: of %d actions performed in both applications, %s was slower in %d, '
                   '%s was slower in %d, and %d were similar.' %
                   (len(compared), n2, s2, n1, s1, len(compared) - s1 - s2))
    else:
        verdict = 'Verdict: no action was timed in both applications, so no direct comparison is possible.'
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    cell = ws.cell(row=row, column=1, value=verdict)
    cell.font = Font(bold=True, size=12, color=NAVY)
    cell.fill = _fill('DEEAF6')
    cell.alignment = Alignment(wrap_text=True, vertical='center', indent=1)
    ws.row_dimensions[row].height = 34

    row = _section(ws, row + 2, 'Actions performed in only one application')
    only = [[c['action'], c['page'], c['status'].replace('Only in ', ''), '', c['note']]
            for c in comparison if c['status'].startswith('Only in')]
    if only:
        start = row
        row = _table(ws, row, ['Action', 'Page', 'Performed in', '', 'Note'], only)
        for r in range(start + 1, row + 1):
            ws.cell(row=r, column=5).fill = _fill('FFF2CC')
    else:
        ws.cell(row=row, column=1, value='None - every action was performed in both applications.')
        row += 1

    row = _section(ws, row + 2, 'How to read this report')
    notes = [
        'Times are measured in the browser: from the click until the last network response finished.',
        'Same page + same button/link label = same action. Differences under 10%% (or %d ms) count as similar.'
        % SIMILAR_MS,
        'Query strings and URL fragments are never recorded. See the Comparison, Actions and Requests sheets for detail.',
    ]
    for note in notes:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        cell = ws.cell(row=row, column=1, value='-  ' + note)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[row].height = 20
        row += 1

    # ---- Comparison --------------------------------------------------------
    headers = ['Action', 'Page', '%s runs' % n1, '%s median (ms)' % n1, '%s runs' % n2,
               '%s median (ms)' % n2, 'Difference (ms)', 'Difference %', 'Result', 'Status', 'Note']
    rows = [[c['action'], c['page'], c['runs1'], c['med1'], c['runs2'], c['med2'], c['diff'], c['pct'],
             c['verdict'], c['status'], c['note']] for c in comparison]
    ws = _data_sheet(wb, 'Comparison', 'Action-by-action comparison',
                     'Positive difference = %s slower than %s.' % (n2, n1),
                     headers, rows, [34, 30, 12, 15, 12, 15, 14, 12, 26, 26, 46], (3, 4, 5, 6, 7))
    for r in range(4, 4 + len(rows)):
        ws.cell(row=r, column=8).number_format = '+0%;-0%;0%'
        result = ws.cell(row=r, column=9)
        if result.value and result.value.endswith('slower'):
            result.fill = _fill('F8CBAD')
            result.font = Font(bold=True, color='9C0006')
        elif result.value == 'Similar':
            result.fill = _fill('E2EFDA')
        if str(ws.cell(row=r, column=10).value).startswith('Only in'):
            for c in (10, 11):
                ws.cell(row=r, column=c).fill = _fill('FFF2CC')
            ws.cell(row=r, column=10).font = Font(bold=True, color='9C5700')

    # ---- Actions -----------------------------------------------------------
    headers = ['#', 'Time', 'Tab', 'Application', 'Page', 'Action', 'Type', 'Outcome', 'Total (ms)',
               'DOM ready (ms)', 'Page load (ms)', 'Requests', 'Failed', 'Slowest request (ms)',
               'Server wait (ms)', 'Slowest request', 'Data (KB)', 'Done in other app?']
    rows = []
    for r in results:
        other = ''
        if r['application'] in (n1, n2):
            other = 'Yes' if (r['page_rel'], r['action'].strip().lower()) in in_both else 'No'
        rows.append([r['action_id'], r['started'], r['tab'], r['application'], r['page_rel'], r['action'],
                     r['action_type'], r['outcome'], r['total_ms'], r['dom_ms'], r['load_ms'],
                     r['request_count'], r['failed_requests'], r['slowest_ms'], r['server_wait_ms'],
                     short_url(r['slowest_url']) if r['slowest_url'] else '', r['kb'], other])
    ws = _data_sheet(wb, 'Actions', 'Every recorded action',
                     'One row per click or page navigation, in the order performed.', headers, rows,
                     [6, 19, 8, 18, 28, 30, 11, 22, 11, 11, 11, 10, 8, 12, 12, 28, 10, 12],
                     (9, 10, 11, 14, 15))
    for r in range(4, 4 + len(rows)):
        total = ws.cell(row=r, column=9)
        if isinstance(total.value, int) and total.value >= slow_ms:
            total.font = Font(bold=True, color='C00000')
        if ws.cell(row=r, column=18).value == 'No':
            ws.cell(row=r, column=18).fill = _fill('FFF2CC')
        if ws.cell(row=r, column=13).value:
            ws.cell(row=r, column=13).font = Font(bold=True, color='C00000')

    # ---- Requests ----------------------------------------------------------
    headers = ['#', 'Application', 'Action', 'Method', 'Type', 'URL (no query)', 'Status', 'Total (ms)',
               'Server wait (ms)', 'KB', 'Error']
    rows = []
    for result, requests in report.completed:
        for q in requests:
            rows.append([result['action_id'], result['application'], result['action'], q['method'],
                         q['type'], q['url'], q['status'], q['duration_ms'], q['ttfb_ms'],
                         round(float(q['bytes_received'] or 0) / 1024, 1), q['error']])
    ws = _data_sheet(wb, 'Requests', 'Network request detail',
                     'Server wait = time until the server sent the first byte of the response.',
                     headers, rows, [6, 18, 28, 9, 12, 60, 8, 11, 14, 9, 26], (8, 9))
    for r in range(4, 4 + len(rows)):
        status = ws.cell(row=r, column=7).value
        if (isinstance(status, int) and status >= 400) or ws.cell(row=r, column=11).value:
            for c in (7, 11):
                ws.cell(row=r, column=c).font = Font(bold=True, color='C00000')

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path


def main():
    parser = argparse.ArgumentParser(description='Compare two web applications side by side in Edge and write an Excel report.')
    parser.add_argument('--app1', default='', help='First application URL (opened in tab 1)')
    parser.add_argument('--app2', default='', help='Second application URL (opened in tab 2)')
    parser.add_argument('--name1', default='Application 1', help='Display name for the first application')
    parser.add_argument('--name2', default='Application 2', help='Display name for the second application')
    parser.add_argument('--slow-ms', type=int, default=DEFAULT_SLOW_MS,
                        help='Actions at or above this many ms are highlighted as slow (default %d)' % DEFAULT_SLOW_MS)
    parser.add_argument('--out', default='', help='Folder for the Excel report (default: current folder)')
    args = parser.parse_args()
    name1, name2 = args.name1.strip() or 'Application 1', args.name2.strip() or 'Application 2'
    if name1.lower() == name2.lower():
        print('--name1 and --name2 must be different.', file=sys.stderr)
        return 2
    if sys.platform != 'win32':
        print('This recorder launches Microsoft Edge on Windows. Running OS: %s' % sys.platform, file=sys.stderr)
        return 2
    if Workbook is None:
        print('openpyxl is required for the Excel report. Install it with:\n    py -3.9 -m pip install openpyxl',
              file=sys.stderr)
        return 2
    process = None
    browser = None
    report = Report(args.app1, args.app2, name1, name2, args.slow_ms)
    profile = tempfile.mkdtemp(prefix='edge_perf_profile_')
    output = (Path(args.out) if args.out else Path.cwd()) / (
        'url_compare_report_' + dt.datetime.now().strftime('%Y%m%d_%H%M%S') + '.xlsx')
    print('Opening a separate Edge browser profile for this test.')
    try:
        edge = find_edge()
        process = subprocess.Popen([
            edge, '--remote-debugging-port=0', '--remote-debugging-address=127.0.0.1',
            '--user-data-dir=' + profile, '--no-first-run', '--no-default-browser-check',
            '--new-window', 'about:blank'
        ], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        port = wait_for_debugging_port(profile, process)
        browser = Browser(port, report)
        pages = browser.pages()
        if not pages:
            raise RuntimeError('No debuggable Edge page found.')
        pages.sort(key=lambda p: p.get('url') != 'about:blank')
        browser.attach(pages[0], 'Tab 1', name1 if args.app1 else None, args.app1 or None)
        for extra in pages[1:]:
            browser.known.add(extra['id'])     # Edge start-up pages, if any
        if args.app2:
            browser.attach(browser.new_tab(), 'Tab 2', name2, args.app2)
        print('Recording. Tab 1 = %s, Tab 2 = %s.' % (name1, name2 if args.app2 else '(none)'))
        print('Work in either tab, in any order. New tabs are recorded automatically.')
        print('Press Ctrl+C here (or close the Edge window) to stop and save the Excel report.\n')
        last_poll = time.monotonic()
        while True:
            browser.pump()
            if time.monotonic() - last_poll >= TAB_POLL_SECONDS:
                last_poll = time.monotonic()
                if process.poll() is not None:
                    print('\nEdge was closed.')
                    break
                browser.sync_tabs()
    except KeyboardInterrupt:
        print('\nStopped by user.')
    except (RuntimeError, OSError, ConnectionError, ValueError) as ex:
        print('\nERROR: %s' % ex, file=sys.stderr)
    finally:
        if browser:
            for tab in list(browser.tabs):
                if tab.active:
                    tab.finish_action('STOPPED_BEFORE_COMPLETION')
                try:
                    tab.ws.close()
                except OSError:
                    pass
        if report.completed:
            try:
                path = write_excel(report, output)
                print('\nSaved %d actions. Excel report:\n  %s' % (len(report.completed), path))
            except OSError as ex:
                print('\nCould not save the Excel report (%s). Is a file with that name open?' % ex,
                      file=sys.stderr)
        else:
            print('\nNo actions were recorded, so no report was written.')
        if process:
            # Terminate only the dedicated Edge process we started, not existing normal Edge sessions.
            try:
                process.terminate()
                process.wait(timeout=8)
            except (OSError, subprocess.TimeoutExpired):
                pass
        try:
            shutil.rmtree(profile)
        except OSError:
            print('Temporary Edge profile could not be deleted; close test Edge and delete: %s' % profile)
    return 0


if __name__ == '__main__':
    sys.exit(main())
